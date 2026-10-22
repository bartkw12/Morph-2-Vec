"""
Feasibility step 4: spreadsheets and notes for the linguistic (theory) annotation of the past-tense
items. Reads outputs/feasibility/data/ (step 3) and the analogy inputs; writes outputs/feasibility/theory/:

    vbd_pairs_for_annotation.xlsx (+ .csv)   one row per past-tense pair, with a class dropdown
    vbd_false_friends_for_review.xlsx        -ed false friends (bold/bowl)
    forced_choice_items_vbd.xlsx             D/T, D/ɪd, T/ɪd forced-choice triples
    hypotheses_draft.md, README_annotation.md

    python scripts/feasibility/03_prepare_theory_materials.py
"""

from argparse import ArgumentParser
from pathlib import Path
import pickle
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C  # noqa: E402

sys.path.insert(0, str(C.REPO_ROOT / "scripts"))
from summarize_analogy_results import FF_EXCLUDE, FF_EXCLUDE_INFLECTED  # noqa: E402

CLASSES = ["regular_t", "regular_d", "regular_id", "semi_regular", "devoicing", "ablaut", "suppletive", "other", "exclude"]
ANNOTATE_FILL = PatternFill("solid", fgColor="FFF2CC")


# --------------------------------------------------------------------------- spreadsheet formatting

def write_xlsx(path: Path, sheets: dict[str, pd.DataFrame], annotate_cols: list[str], dropdowns: dict[str, list[str]]):
    """Write sheets, then: bold frozen header, auto-sized columns, shaded columns to fill in, dropdowns."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    for name, df in sheets.items():
        ws = wb[name]
        ws.freeze_panes = "A2"
        for j, col in enumerate(df.columns, start=1):
            letter = get_column_letter(j)
            ws[f"{letter}1"].font = Font(bold=True)
            width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
            ws.column_dimensions[letter].width = min(max(8, width + 2), 50)
            last = max(len(df) + 1, 2)
            if col in annotate_cols:
                for row in range(1, last + 1):
                    ws[f"{letter}{row}"].fill = ANNOTATE_FILL
            if col in dropdowns:
                dv = DataValidation(type="list", formula1='"' + ",".join(dropdowns[col]) + '"', allow_blank=True,
                                    showErrorMessage=True, errorTitle="Not in list",
                                    error="Pick a value from the dropdown (or leave it empty).")
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{last}")
    wb.save(path)


# --------------------------------------------------------------------------- 1. VBD pairs

def pairs_sheet(pos_counts) -> pd.DataFrame:
    p = pd.read_parquet(C.DATA / "pairs_vbd.parquet")
    freq = {w: sum(c.values()) for w, c in pos_counts.items()}
    p["base_corpus_count"] = p.base.map(freq).fillna(0).astype(int)
    p["past_corpus_count"] = p.inflected.map(freq).fillna(0).astype(int)
    guesses = p.apply(C.auto_class, axis=1)
    p["auto_class"], p["auto_reason"] = guesses.str[0], guesses.str[1]
    p = p.sort_values(["is_regular", "past_corpus_count"], ascending=[True, False])
    return pd.DataFrame({
        "base": p.base, "past": p.inflected,
        "base_tokens": p.base_tokens, "past_tokens": p.past_tokens,
        "base_corpus_count": p.base_corpus_count, "past_corpus_count": p.past_corpus_count,
        "base_pron": p.base_pron, "base_ipa": p.base_pron.map(C.to_ipa),
        "past_pron": p.past_pron, "past_ipa": p.past_pron.map(C.to_ipa),
        "ending": p.post_divergence, "ending_ipa": p.post_divergence.map(C.to_ipa),
        "share_ending_AH_D": p.share_AH_D.round(2),
        "auto_regular": p.is_regular, "auto_class": p.auto_class, "auto_reason": p.auto_reason,
        "authors_excluded": p.authors_excluded,
        "past_sounds_like": p.past_homophones, "n_pasts_for_base": p.n_pasts_for_base,
        "class": "", "sound_change_family": "", "exclude": "", "notes": "",
    })


# --------------------------------------------------------------------------- 2. false friends

def false_friends_sheet(counts: pd.Series) -> pd.DataFrame:
    ff = pd.read_csv(C.INPUTS / "false_friends.csv")
    ff = ff[ff.inflection == "VBD"].copy()
    override = [C.VBD_STRONG_OVERRIDES.get((b, i, d)) for b, i, d in zip(ff.base, ff.inflected, ff.post_divergence)]
    ff["authors_strong_override"] = [("" if o is None else o) for o in override]
    ff["strong_used"] = [s if o is None else o for s, o in zip(ff.strong, override)]
    ff["authors_excluded"] = ff.base.isin(FF_EXCLUDE) | ff.inflected.isin(FF_EXCLUDE_INFLECTED)
    n_bases = ff.groupby("post_divergence").base.transform("nunique")
    ff["in_analysis"] = ff.strong_used.astype(bool) & ~ff.authors_excluded & (n_bases >= 10)
    ff = ff.sort_values(["post_divergence", "base"])
    return pd.DataFrame({
        "base": ff.base, "false_friend": ff.inflected,
        "base_tokens": ff.base.map(counts), "false_friend_tokens": ff.inflected.map(counts),
        "base_pron": ff.base_form, "base_ipa": ff.base_form.map(C.to_ipa),
        "false_friend_pron": ff.inflected_form, "false_friend_ipa": ff.inflected_form.map(C.to_ipa),
        "added_ending": ff.post_divergence.map(C.to_ipa), "rule_predicts": ff.strong_expected.map(C.to_ipa),
        "pipeline_strong": ff.strong, "authors_strong_override": ff.authors_strong_override,
        "authors_excluded": ff.authors_excluded, "in_analysis": ff.in_analysis,
        "keep": "", "notes": "",
    })


# --------------------------------------------------------------------------- 3. forced choice

def forced_choice_sheet() -> pd.DataFrame:
    sets = ["FC-D_T", "FC-D_IH D", "FC-T_IH D"]
    aci = pd.read_parquet(C.INPUTS / "all_cross_instances.parquet",
                          columns=["inflection", "base", "base_phones", "inflected", "inflected_phones",
                                   "inflected2", "inflected2_phones"],
                          filters=[("inflection", "in", sets)])
    items = aci.groupby(["inflection", "base", "inflected", "inflected2"]).head(1).copy()
    used = pd.read_parquet(C.DATA / "fc_trials.parquet", columns=["inflection_to", "base_to", "inflected_to"]) \
        .drop_duplicates().assign(in_analysis=True)
    items = items.merge(used.rename(columns={"inflection_to": "inflection", "base_to": "base", "inflected_to": "inflected"}),
                        how="left", on=["inflection", "base", "inflected"])
    items["in_analysis"] = items.in_analysis.fillna(False).astype(bool)

    rows = []
    for r in items.itertuples():
        expected = C.guess_past_allomorph(r.base_phones.split())
        suf1 = r.inflected_phones[len(r.base_phones) + 1:]
        suf2 = r.inflected2_phones[len(r.base_phones) + 1:]
        forms = [(r.inflected, r.inflected_phones, suf1), (r.inflected2, r.inflected2_phones, suf2)]
        if suf2 == expected:
            forms.reverse()
        consistent = forms[0] if forms[0][2] == expected else None
        rows.append({
            "set": {"FC-D_T": "d / t", "FC-D_IH D": "d / ɪd", "FC-T_IH D": "t / ɪd"}[r.inflection],
            "base": r.base, "base_pron": r.base_phones, "base_ipa": C.to_ipa(r.base_phones),
            "rule_predicts": C.to_ipa(expected),
            "consistent_form": forms[0][0] if consistent else "", "consistent_pron": C.to_ipa(forms[0][1]) if consistent else "",
            "inconsistent_form": forms[1][0] if consistent else "", "inconsistent_pron": C.to_ipa(forms[1][1]) if consistent else "",
            "form_1": "" if consistent else r.inflected, "form_1_pron": "" if consistent else C.to_ipa(r.inflected_phones),
            "form_2": "" if consistent else r.inflected2, "form_2_pron": "" if consistent else C.to_ipa(r.inflected2_phones),
            "in_analysis": r.in_analysis, "ok": "", "notes": "",
        })
    return pd.DataFrame(rows).sort_values(["set", "base"])


# --------------------------------------------------------------------------- 4-5. notes

HYPOTHESES = """# Hypotheses draft: past tense in a word-recognition probe

A draft for the theory review to refine. It feeds the proposal section "hypothesis and what each outcome would mean".

What we can measure (see `README_annotation.md`):
- **Offset similarity.** Is the change *walk → walked* the same direction in the model's space as *play → played*, *want → wanted*, *keep → kept*, *sing → sang*?
- **Transfer.** Does the change learned from one verb predict another verb's past form?
- **Held-out verbs.** Does that work for verbs the average was not computed from?

| Account | Regular [t]/[d]/[ɪd] | Semi-regular (kept, slept) | Irregular (sang, took) | -ed false friends (bold/bowl) |
|---|---|---|---|---|
| Dual-route (Pinker) | one shared offset across allomorphs | partly shared | word-specific, no transfer | should not share the offset |
| Single-route (Rumelhart & McClelland) | shared, graded by phonology | intermediate | clusters by sound-change family (sing/ring/drink) | partly shared |
| Surface phonology (Gauthier et al. finding for -s) | shared | shares the regular [t] offset | no transfer | shares the offset |

Open points:
- Is "shares the regular [t] offset" the right prediction for semi-regulars under the surface-phonology account, or should *kept* pattern with [t] only in its final segment?
- Which sound-change families have enough verbs in our data to test the single-route clustering prediction? (See the `sound_change_family` column.)
- What should we predict for devoicing verbs (*send/sent, build/built*)?
- What would count as "partly shared", in numbers? For example, cosine above the random-pair control but below the regular-regular level.
"""

README = """# Annotation guide: past-tense items

These files are for the linguistic annotation of the past-tense verbs in our data. No code is needed: open them in Excel.
**The yellow columns are the ones to fill in**; everything else was filled in automatically.
A first pass of **2–3 hours on the top ~150 rows of the first file is enough for now.**

## Background in two sentences

The model turns each spoken word into a point in a 32-dimensional space. We test whether "make it past tense"
is one consistent direction in that space (*walk → walked* ≈ *play → played* ≈ *want → wanted*), and how
semi-regular and irregular verbs relate to it. The annotated classes decide which verbs are grouped together in that test.

## The files

### 1. `vbd_pairs_for_annotation.xlsx` (main task)

There is one row per verb pair found in the audio: 949 pairs, irregulars first, then by how common the past form is.

| Column | Meaning |
|---|---|
| base, past | the two word forms |
| base_tokens, past_tokens | how many spoken examples the model sees (capped at 100) |
| base_corpus_count, past_corpus_count | how often the word occurs in the whole 100-hour corpus |
| base_pron / past_pron | the pronunciation the speech aligner used most often, in ARPAbet (the CMU dictionary alphabet) |
| base_ipa / past_ipa | the same, converted to IPA. Stress marks were removed, so **ʌ also stands for schwa** |
| ending | what the past form adds after the point where it stops sounding like the base |
| share_ending_AH_D | for -ed syllables: the share of examples the aligner wrote as [ʌd] rather than [ɪd] (both mean ɪd) |
| auto_regular | the program's spelling-based guess: does the past end in -ed/-d/-t added to the base? |
| auto_class, auto_reason | an automatic first guess at the class (below), and why. `check` means "please look at this one" |
| authors_excluded | the original paper's authors left this verb out (stress-shift verbs like *record*, or odd cases like *do/did*) |
| past_sounds_like | other words in the data with the same pronunciation as the past form (*made* ~ *maid*). These make the test harder for that verb |
| n_pasts_for_base | more than 1 means the program found two past forms (*spring → sprang, sprung*) |
| **class** | the class, from the dropdown |
| **sound_change_family** | free text, e.g. `sing-sang`, `keep-kept`, `bring-brought`; use the same label for verbs that change the same way |
| **exclude** | `yes` if the pair should not be used at all (not really a past tense, a different word, an alignment error) |
| **notes** | anything else |

**Classes**, with two examples each:

| class | meaning | examples |
|---|---|---|
| regular_t | regular, ending pronounced [t] | walked, missed |
| regular_d | regular, ending pronounced [d] | played, called |
| regular_id | regular, ending pronounced [ɪd] | wanted, needed |
| semi_regular | adds [t]/[d] *and* changes the stem vowel | kept, slept |
| devoicing | final [d] of the base becomes [t] | sent, built |
| ablaut | vowel change only, no added ending | sang, took |
| suppletive | a different root | went, was |
| other | anything that fits none of these (*brought*? *made*?), with a note | — |
| exclude | not a usable past-tense pair | could (from *can*), wound (from *wind*) |

### 2. `vbd_false_friends_for_review.xlsx`

These are word pairs that *sound* like a verb and its past but aren't one, e.g. *bold* ~ *bowl* + [d], *wild* ~ *wile* + [d], *plant* ~ *plan* + [t].
The model's -s results in the original paper depend on such "false friends", so we need a clean list for -ed.
- `pipeline_strong` / `authors_strong_override`: whether the added sound is the one the regular rule would predict (*cold* = *coal* + [d] after a voiced sound: yes).
- `in_analysis`: whether the pair is currently used. The [ɪd] set is too small to use (8 pairs).
- **keep**: `yes`/`no`. Is this a genuine look-alike, i.e. not secretly a real inflection and not a name or foreign word? **notes**: why.

### 3. `forced_choice_items_vbd.xlsx`

These are triples where both a [d] and a [t] (or [ɪd]) version of a base exist as words, e.g. *coal → cold / colt*.
The test asks which one the model prefers when we apply a past-tense change to the base.
The `consistent_form` is the one the regular voicing rule predicts.
- Please check the pronunciations and say whether the triple makes sense (**ok**: yes/no/unsure).
- Rows with `form_1` / `form_2` filled are cases where neither form matches the rule's prediction.

## Hypotheses

`hypotheses_draft.md` has the table of predictions per theory. Please refine the wording and the predictions,
especially for semi-regular and devoicing verbs.

Please don't rename the columns: the scripts read the completed files back by column name.
"""


def main(args):
    C.THEORY.mkdir(parents=True, exist_ok=True)
    with open(C.POS_COUNTS, "rb") as f:
        pos_counts = pickle.load(f)
    counts = C.load_spec().label_counts

    pairs = pairs_sheet(pos_counts)
    pairs.to_csv(C.THEORY / "vbd_pairs_for_annotation.csv", index=False, encoding="utf-8-sig")
    write_xlsx(C.THEORY / "vbd_pairs_for_annotation.xlsx", {"pairs": pairs},
               annotate_cols=["class", "sound_change_family", "exclude", "notes"],
               dropdowns={"class": CLASSES, "exclude": ["yes", "no"]})
    print(f"vbd_pairs_for_annotation: {len(pairs)} rows; auto_class: {pairs.auto_class.value_counts().to_dict()}")

    ff = false_friends_sheet(counts)
    write_xlsx(C.THEORY / "vbd_false_friends_for_review.xlsx", {"false_friends": ff},
               annotate_cols=["keep", "notes"], dropdowns={"keep": ["yes", "no"]})
    print(f"vbd_false_friends_for_review: {len(ff)} rows, {int(ff.in_analysis.sum())} in the analysis")

    fc = forced_choice_sheet()
    write_xlsx(C.THEORY / "forced_choice_items_vbd.xlsx", {"forced_choice": fc},
               annotate_cols=["ok", "notes"], dropdowns={"ok": ["yes", "no", "unsure"]})
    print(f"forced_choice_items_vbd: {fc.groupby('set').size().to_dict()}, {int(fc.in_analysis.sum())} in the analysis")

    (C.THEORY / "hypotheses_draft.md").write_text(HYPOTHESES, encoding="utf-8")
    (C.THEORY / "README_annotation.md").write_text(README, encoding="utf-8")
    print("wrote", C.THEORY)


if __name__ == "__main__":
    p = ArgumentParser(description=__doc__)
    main(p.parse_args())
